# FU87 — ekspor XLSX Ketersediaan Stock: JEMBATAN baris tampilan halaman
# (mock FU83) → workbook 3 sheet. _siapkan_sheet/_nama_file_xlsx pure
# (tanpa DB — plain unittest); export_xlsx dibuktikan menghasilkan binary
# zip dgn nama file paritas helper frontend (pola test FU80c di test_dashboard).

import unittest
from io import BytesIO

import frappe

from production_app.api.stock_availability import (
	_nama_file_xlsx,
	_siapkan_sheet,
	export_xlsx,
)

# FU88: klien mengonversi qty ke UOM tampilan (DIU) sebelum mengirim —
# stok_diu/reserved_diu/tersedia_diu/minimum_diu + uom_tampilan; actual_qty
# tetap stock UOM utk kolom "Stok (Stock UOM)".
SAMPLE_ITEMS = [
	{
		"item_code": "RM-FLOUR-001",
		"item_name": "Tepung Terigu",
		"item_group": "Bahan Baku",
		"actual_qty": 125,
		"stock_uom": "kg",
		"uom_tampilan": "kg",
		"stok_diu": 125,
		"reserved_diu": 20,
		"tersedia_diu": 105,
		"minimum_diu": 100,
		"status": "aman",
		"capacity_batch": 4,
		"capacity_detail": [
			{"product": "Dough Coklat", "batches": 4},
			{"product": "Dough Keju", "batches": 3},
		],
	},
	{
		"item_code": "RM-YEAST-006",
		"item_name": "Ragi Instan",
		"item_group": "Bahan Baku",
		"actual_qty": 22,
		"stock_uom": "kg",
		"uom_tampilan": "Pack",
		"stok_diu": 5.5,
		"reserved_diu": 1,
		"tersedia_diu": 4.5,
		"minimum_diu": 3.75,
		"status": "aman",
		"capacity_batch": 8,
		"capacity_detail": [{"product": "Dough Coklat", "batches": 8}],
	},
	{
		"item_code": "RM-COCOA-004",
		"item_name": "Bubuk Kakao",
		"item_group": "Bahan Baku",
		"actual_qty": 18,
		"stock_uom": "kg",
		"uom_tampilan": "kg",
		"stok_diu": 18,
		"reserved_diu": 18,
		"tersedia_diu": 0,
		"minimum_diu": 25,
		"status": "habis",
		"capacity_batch": 0,
		"capacity_detail": [{"product": "Dough Coklat", "batches": 0}],
	},
]


class TestStockAvailabilitySheet(unittest.TestCase):
	"""Pure sheet-building — tanpa fixture DB."""

	def test_siapkan_sheet_tiga_sheet_kontrak(self):
		info, baris, kapasitas = _siapkan_sheet(
			SAMPLE_ITEMS, {"aman": 96, "menipis": 21, "habis": 11}, "Gudang Produksi"
		)
		# Info: judul + ringkasan KPI + gudang
		self.assertEqual(info[0][0], "Ekspor Ketersediaan Stock Bahan Baku")
		label = {r[0]: r[1] for r in info[1:]}
		self.assertEqual(label["Gudang"], "Gudang Produksi")
		self.assertEqual(label["Jumlah Item"], 3)
		self.assertEqual((label["Stock Aman"], label["Stock Menipis"], label["Stock Habis"]), (96, 21, 11))
		self.assertIn("Dicetak", label)
		# Stock Bahan Baku: header 11 kolom (FU88: qty = UOM tampilan + kolom
		# penutup Stok (Stock UOM)); angka MENTAH numerik
		self.assertEqual(
			baris[0],
			["Kode Item", "Nama Item", "Grup Item", "Stok", "Reserved", "Tersedia",
			 "UOM", "Stok Minimum", "Status", "Kapasitas (Batch)", "Stok (Stock UOM)"],
		)
		tepung = baris[1]  # tanpa DIU → diu = stock, UOM = stock
		self.assertEqual(tepung[0], "RM-FLOUR-001")
		self.assertIsInstance(tepung[3], float)  # stok numerik (flt)
		self.assertEqual(tepung[5], 105.0)  # tersedia = stok − reserved
		self.assertEqual(tepung[6], "kg")
		self.assertEqual(tepung[8], "Aman")  # label Indonesia dari status angka
		self.assertEqual(tepung[10], 125.0)  # stok mentah stock UOM
		ragi = baris[2]  # ber-DIU: qty dalam Pack, stok mentah tetap kg
		self.assertEqual(ragi[3], 5.5)
		self.assertEqual(ragi[6], "Pack")
		self.assertEqual(ragi[10], 22.0)
		kakao = baris[3]
		self.assertEqual(kakao[8], "Habis")
		self.assertEqual(kakao[9], 0)
		# Kapasitas per Resep: diratakan 1 baris per resep (4 baris data)
		self.assertEqual(
			kapasitas[0],
			["Kode Item", "Nama Item", "Produk Resep", "Kapasitas (Batch)"],
		)
		self.assertEqual(len(kapasitas), 5)
		self.assertEqual(kapasitas[1][2], "Dough Coklat")
		self.assertEqual(kapasitas[1][3], 4)
		self.assertEqual(kapasitas[2][2], "Dough Keju")  # resep kedua Tepung
		self.assertEqual(kapasitas[4][2], "Dough Coklat")  # satu-satunya resep Kakao

	def test_siapkan_sheet_kosong_tahan(self):
		info, baris, kapasitas = _siapkan_sheet([], {}, None)
		self.assertEqual(len(baris), 1)  # hanya header
		self.assertEqual(len(kapasitas), 1)
		self.assertEqual(info[1][1], "Gudang Produksi")  # default tanpa gudang

	def test_siapkan_sheet_fallback_tanpa_uom_tampilan(self):
		# item lama tanpa kunci FU88 → UOM fallback stock_uom, diu 0.0 (flt)
		_, baris, _ = _siapkan_sheet(
			[{"item_code": "RM-LAMA", "stock_uom": "kg", "status": "aman"}], {}, None
		)
		self.assertEqual(baris[1][6], "kg")
		self.assertEqual(baris[1][10], 0.0)

	def test_siapkan_sheet_minimum_kapasitas_kosong_bukan_nol(self):
		# P2-2 review FU90: minimum/kapasitas None → sel KOSONG (paritas layar
		# '-'); 0 asli tetap 0 — None vs 0 dibedakan
		_, baris, _ = _siapkan_sheet(
			[
				{
					"item_code": "RM-KOSONG",
					"item_name": "Tanpa Minimum",
					"item_group": "Bahan Baku",
					"stock_uom": "kg",
					"uom_tampilan": "kg",
					"actual_qty": 5,
					"stok_diu": 5,
					"reserved_diu": 0,
					"tersedia_diu": 5,
					"minimum_diu": None,
					"status": "aman",
					"capacity_batch": None,
					"capacity_detail": [],
				}
			],
			{},
			None,
		)
		self.assertEqual(baris[1][7], "")  # Stok Minimum kosong (bukan 0)
		self.assertEqual(baris[1][9], "")  # Kapasitas (Batch) kosong (bukan 0)
		# 0 asli tetap 0
		_, baris_nol, _ = _siapkan_sheet(
			[{"item_code": "RM-NOL", "stock_uom": "kg", "minimum_diu": 0, "capacity_batch": 0}],
			{},
			None,
		)
		self.assertEqual(baris_nol[1][7], 0.0)
		self.assertEqual(baris_nol[1][9], 0)


class TestStockAvailabilityExport(unittest.TestCase):
	def test_nama_file_paritas_frontend(self):
		self.assertEqual(
			_nama_file_xlsx("Gudang Produksi"),
			f"ketersediaan-stock-gudang-produksi-{frappe.utils.nowdate()}.xlsx",
		)
		self.assertEqual(
			_nama_file_xlsx(None), f"ketersediaan-stock-{frappe.utils.nowdate()}.xlsx"
		)

	def test_export_xlsx_binary_tiga_sheet(self):
		"""Endpoint men-set frappe.response binary (provide_binary_file):
		filecontent bytes mulai magic PK zip, nama .xlsx paritas frontend."""
		export_xlsx(items=SAMPLE_ITEMS, summary={"aman": 96, "menipis": 21, "habis": 11}, warehouse="Gudang Produksi")
		self.assertEqual(frappe.response["type"], "binary")
		content = frappe.response["filecontent"]
		self.assertIsInstance(content, bytes)
		self.assertTrue(content.startswith(b"PK"))  # magic zip .xlsx
		self.assertTrue(frappe.response["filename"].startswith("ketersediaan-stock-gudang-produksi"))
		self.assertTrue(frappe.response["filename"].endswith(".xlsx"))
		# workbook valid: 3 sheet sesuai kontrak (buka balik via openpyxl)
		from openpyxl import load_workbook

		wb = load_workbook(BytesIO(content))
		self.assertEqual(wb.sheetnames, ["Info", "Stock Bahan Baku", "Kapasitas per Resep"])
		self.assertEqual(wb["Stock Bahan Baku"].max_row, 4)  # header + 3 item


# ------------------------------------------------------------------ FU90
# Integrasi ERPNext nyata: helper murni diuji tanpa DB; kontrak endpoint
# diuji live (run-tests = Administrator) TANPA asersi angka eksak data live
# (delta-safe — data berubah kapan pun, hanya bentuk/tipe/konsistensi
# internal yang dipatok).

from datetime import date, datetime

from production_app.api.stock_availability import (
	_kapasitas_dari_bom,
	_min_stock_item,
	_peta_movement,
	_status_stock,
	stock_availability,
	stock_movements,
)


def _item_stub(stock_uom="kg", uoms=()):
	"""Item minimal utk helper murni (kunlinkan _conversion_factor)."""
	return frappe._dict(stock_uom=stock_uom, variant_of=None, uoms=list(uoms))


class TestStockAvailabilityHelpers(unittest.TestCase):
	"""Helper murni FU90 — tanpa DB."""

	def test_status_batas(self):
		self.assertEqual(_status_stock(0, None), "habis")
		self.assertEqual(_status_stock(-0.5, None), "habis")
		self.assertEqual(_status_stock(10, None), "aman")  # minimum kosong tak pernah menipis
		self.assertEqual(_status_stock(10, 10.0), "aman")  # tepat sama dgn minimum bukan menipis
		self.assertEqual(_status_stock(9.999, 10.0), "menipis")
		self.assertEqual(_status_stock(5, 0.0), "aman")

	def test_min_stock_gudang_dan_konversi(self):
		rows = [
			frappe._dict(warehouse="Gudang Lain", warehouse_reorder_level=99.0),
			frappe._dict(warehouse="W", warehouse_reorder_level=40.0),
		]
		item = _item_stub()
		self.assertEqual(_min_stock_item(item, rows, "W"), 40.0)  # baris gudang lain diabaikan
		self.assertIsNone(_min_stock_item(item, rows, "Tidak Ada"))
		self.assertIsNone(_min_stock_item(item, [], "W"))
		# warehouse_uom sama dengan stock_uom → tanpa konversi
		sama = frappe._dict(warehouse="W", warehouse_reorder_level=3.0, warehouse_uom="kg")
		self.assertEqual(_min_stock_item(item, [sama], "W"), 3.0)
		# warehouse_uom beda → konversi faktor item-specific (2 Sak × 25 kg)
		item_sak = _item_stub(uoms=[frappe._dict(uom="Sak", conversion_factor=25.0)])
		baris_sak = frappe._dict(warehouse="W", warehouse_reorder_level=2.0, warehouse_uom="Sak")
		self.assertEqual(_min_stock_item(item_sak, [baris_sak], "W"), 50.0)
		# faktor tak ditemukan/invalid → None (jangan mengarang)
		self.assertIsNone(_min_stock_item(item, [baris_sak], "W"))
		faktor_nol = frappe._dict(warehouse="W", warehouse_reorder_level=2.0, warehouse_uom="Sak")
		item_nol = _item_stub(uoms=[frappe._dict(uom="Sak", conversion_factor=0.0)])
		self.assertIsNone(_min_stock_item(item_nol, [faktor_nol], "W"))

	def test_kapasitas_floor_min_dan_max(self):
		# floor per baris lalu min antar baris: 10.5/3=3.5→3, 7/2=3.5→3
		hasil = _kapasitas_dari_bom(
			{"A": 10.5, "B": 7.0}, [("Roti Coklat", [("A", 3.0), ("B", 2.0)])]
		)
		self.assertEqual(hasil["A"], (3, [{"product": "Roti Coklat", "batches": 3}]))
		self.assertEqual(hasil["B"][0], 3)

	def test_kapasitas_baris_nol_negatif_dan_bahan_luar_gudang(self):
		resep = [("X", [("A", 0.0), ("A", -1.0), ("A", 2.5), ("Z-LUAR", 1000.0)])]
		hasil = _kapasitas_dari_bom({"A": 10.5}, resep)
		self.assertEqual(hasil["A"][0], 4)  # floor(10.5/2.5); need ≤0 di-skip; luar gudang diabaikan
		# FG tanpa bahan gudang ini tidak masuk; item tanpa resep → absen (null/[])
		self.assertEqual(_kapasitas_dari_bom({"A": 10.5}, [("Y", [("Z", 1.0)])]), {})
		self.assertEqual(_kapasitas_dari_bom({"A": 10.5}, []), {})

	def test_kapasitas_negatif_nol_dan_multi_resep(self):
		# tersedia negatif → 0 batch (bukan negatif)
		hasil = _kapasitas_dari_bom({"A": -5.0}, [("X", [("A", 2.0)])])
		self.assertEqual(hasil["A"][0], 0)
		# multi resep: capacity_batch = max antar-FG, detail urut batches desc lalu nama asc
		resep = [("Roti B", [("A", 5.0)]), ("Roti C", [("A", 5.0)]), ("Roti A", [("A", 2.0)])]
		cap, detail = _kapasitas_dari_bom({"A": 10.0}, resep)["A"]
		self.assertEqual(cap, 5)
		self.assertEqual(
			detail,
			[
				{"product": "Roti A", "batches": 5},
				{"product": "Roti B", "batches": 2},
				{"product": "Roti C", "batches": 2},
			],
		)

	def test_peta_movement(self):
		self.assertEqual(
			_peta_movement(
				frappe._dict(
					posting_date=date(2026, 10, 3),
					voucher_type="Material Transfer",
					voucher_no="MAT-STE-2026-00123",
					actual_qty=20.0,
					qty_after_transaction=125.0,
				)
			),
			{
				"tanggal": "2026-10-03",
				"jenis": "Material Transfer",
				"referensi": "MAT-STE-2026-00123",
				"masuk": 20.0,
				"keluar": 0.0,
				"saldo": 125.0,
			},
		)
		keluar = _peta_movement(
			frappe._dict(
				posting_date=date(2026, 10, 2),
				voucher_type="Manufacture",
				voucher_no="X",
				actual_qty=-3.5,
				qty_after_transaction=7.0,
			)
		)
		self.assertEqual((keluar["masuk"], keluar["keluar"], keluar["saldo"]), (0.0, 3.5, 7.0))


class TestStockAvailabilityLive(unittest.TestCase):
	"""Kontrak endpoint di live site. Delta-safe: TANPA asersi angka eksak —
	hanya bentuk respons, tipe field, dan konsistensi internal."""

	KUNCI_ITEM = {
		"item_code",
		"item_name",
		"item_group",
		"stock_uom",
		"actual_qty",
		"reserved_qty",
		"available",
		"display_uom",
		"qty_in_pack",
		"min_stock",
		"status",
		"capacity_batch",
		"capacity_detail",
		"warehouse",
	}

	def test_kontrak_respons_lengkap(self):
		resp = stock_availability()
		self.assertEqual(set(resp), {"warehouse", "warehouses", "generated_at", "summary", "items"})
		self.assertIsInstance(resp["warehouses"], list)
		# urut abjad ala SQL (case-insensitive) — bukan urutan byte Python
		self.assertEqual(resp["warehouses"], sorted(resp["warehouses"], key=str.casefold))
		if resp["warehouse"] is not None:
			self.assertIn(resp["warehouse"], resp["warehouses"])
		datetime.fromisoformat(resp["generated_at"])
		hitung = {"aman": 0, "menipis": 0, "habis": 0}
		for it in resp["items"]:
			self.assertEqual(set(it), self.KUNCI_ITEM)
			self.assertNotIn("movements", it)  # movement lazy — endpoint terpisah
			self.assertIn(it["status"], ("aman", "menipis", "habis"))
			hitung[it["status"]] += 1
			self.assertIsInstance(it["actual_qty"], float)
			self.assertIsInstance(it["reserved_qty"], float)
			self.assertIsInstance(it["available"], float)
			self.assertEqual(it["available"], it["actual_qty"] - it["reserved_qty"])  # server hitung
			self.assertEqual(it["warehouse"], resp["warehouse"])
			if it["display_uom"] is None:
				self.assertIsNone(it["qty_in_pack"])
			else:
				self.assertIsInstance(it["qty_in_pack"], float)
				self.assertGreater(it["qty_in_pack"], 0)
			if it["min_stock"] is not None:
				self.assertIsInstance(it["min_stock"], float)
			self.assertIsInstance(it["capacity_detail"], list)
			if it["capacity_batch"] is None:
				self.assertEqual(it["capacity_detail"], [])  # tanpa resep → null + []
			else:
				self.assertIsInstance(it["capacity_batch"], int)
				self.assertEqual(it["capacity_batch"], max(d["batches"] for d in it["capacity_detail"]))
				urut = [(-d["batches"], d["product"]) for d in it["capacity_detail"]]
				self.assertEqual(urut, sorted(urut))  # batches desc lalu nama asc
			for d in it["capacity_detail"]:
				self.assertEqual(set(d), {"product", "batches"})
				self.assertIsInstance(d["batches"], int)
		self.assertEqual(resp["summary"], {"total": len(resp["items"]), **hitung})

	def test_warehouse_param_dan_gudang_tak_ada(self):
		resp = stock_availability()
		if not resp["warehouses"]:
			self.assertEqual(resp["items"], [])  # site tanpa Bin: kontrak utuh tanpa error
			return
		target = resp["warehouse"] or resp["warehouses"][0]
		resp2 = stock_availability(warehouse=target)
		self.assertEqual(resp2["warehouse"], target)
		self.assertEqual([i["item_code"] for i in resp2["items"]], [i["item_code"] for i in resp["items"]])
		# gudang kosong/tak ada: TIDAK error — items [] + summary nol
		resp3 = stock_availability(warehouse="Gudang Tidak Ada FU90")
		self.assertEqual(resp3["warehouse"], "Gudang Tidak Ada FU90")
		self.assertEqual(resp3["items"], [])
		self.assertEqual(resp3["summary"], {"total": 0, "aman": 0, "menipis": 0, "habis": 0})

	def test_movements_kontrak(self):
		resp = stock_availability()
		if not resp["items"]:
			self.skipTest("site tanpa Bin — tidak ada item utk movement")
		item = resp["items"][0]
		rows = stock_movements(item_code=item["item_code"], warehouse=item["warehouse"])
		self.assertIsInstance(rows, list)
		self.assertLessEqual(len(rows), 10)
		tanggal = [r["tanggal"] for r in rows]
		self.assertEqual(tanggal, sorted(tanggal, reverse=True))  # posting terbaru dulu
		for r in rows:
			self.assertEqual(set(r), {"tanggal", "jenis", "referensi", "masuk", "keluar", "saldo"})
			datetime.strptime(r["tanggal"], "%Y-%m-%d")
			self.assertGreaterEqual(r["masuk"], 0)
			self.assertGreaterEqual(r["keluar"], 0)
			self.assertFalse(r["masuk"] > 0 and r["keluar"] > 0)  # satu sisi dari actual_qty
			self.assertIsInstance(r["saldo"], float)
		# item tanpa SLE → list kosong, bukan error
		self.assertEqual(
			stock_movements(item_code="__FU90_TIDAK_ADA__", warehouse=item["warehouse"]), []
		)
		# item_code kosong → throw ringkas
		with self.assertRaises(frappe.ValidationError):
			stock_movements()
